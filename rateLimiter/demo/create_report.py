import docx
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = docx.Document()

# Adjust margins
sections = doc.sections
for section in sections:
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

# Title
title = doc.add_heading('Design and Implement an API Rate Limiter using Token Bucket', 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

# Header Information Table
table = doc.add_table(rows=2, cols=2)
# We can use a basic style or no style and just set text
# 'Table Grid' is standard. If not found, it won't crash but let's just use it or leave it default.
try:
    table.style = 'Table Grid'
except:
    pass
    
row0 = table.rows[0].cells
row0[0].text = "Date: 01.09.2026"
row0[1].text = "Name: Amudhan A"
row0[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT

row1 = table.rows[1].cells
row1[0].text = "Ex. No: 6"
row1[1].text = "Reg No: 3122 24 5001 016"
row1[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT

# Department Name
dept = doc.add_paragraph('\nDepartment of Computer Science and Engineering')
dept.alignment = WD_ALIGN_PARAGRAPH.CENTER
dept.runs[0].bold = True

# Objective
doc.add_heading('Objective', level=2)
doc.add_paragraph('To implement a Rate Limiter using the Token Bucket algorithm in a Spring Boot application, use Redis to maintain rate-limit information, and analyze request control under normal and burst traffic.')

# Software Used and its Configuration
doc.add_heading('Software Used and its Configuration', level=2)
sw_table = doc.add_table(rows=1, cols=3)
try:
    sw_table.style = 'Table Grid'
except:
    pass
hdr_cells = sw_table.rows[0].cells
hdr_cells[0].text = 'Software'
hdr_cells[1].text = 'Role'
hdr_cells[2].text = 'Purpose in this Assignment'
# make header bold
for cell in hdr_cells:
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.bold = True

data = [
    ('Java & Spring Boot 3+', 'Application Framework', 'Provides the REST API endpoints and web interceptor mechanism to enforce rate limits.'),
    ('Redis', 'In-Memory Data Store', 'Stores the client token counts and timestamps centrally for atomic updates via Lua scripts.'),
    ('Lua Scripting', 'Database Scripting', 'Runs directly inside Redis to ensure rate-limiting state reads and writes are atomic.'),
    ('Docker (Desktop/Compose)', 'Container Runtime', 'Runs the Redis server locally on port 6379 for the Spring Boot app to connect to.'),
    ('Postman / Curl', 'API Testing Tool', 'Generates normal and burst traffic requests to observe rate-limiting behaviour.')
]

for item in data:
    row_cells = sw_table.add_row().cells
    row_cells[0].text = item[0]
    row_cells[1].text = item[1]
    row_cells[2].text = item[2]

# System Design and Architecture
doc.add_heading('System Design and Architecture', level=2)
sys_arch = doc.add_paragraph()
sys_arch.add_run('RateLimitInterceptor: ').bold = True
sys_arch.add_run('Intercepts incoming HTTP requests, extracts the client IP/ID, and checks the rate limit before reaching the controller.\n')
sys_arch.add_run('TokenBucketRateLimiter: ').bold = True
sys_arch.add_run('Service layer that executes the LUA script in Redis to evaluate the Token Bucket algorithm for each request.\n')
sys_arch.add_run('Redis & Lua Script: ').bold = True
sys_arch.add_run('Centralized store that atomically manages token counts and last refill timestamps for every unique client.\n')
sys_arch.add_run('DemoController: ').bold = True
sys_arch.add_run('A sample REST API endpoint (/api/ping) that returns a response if the request successfully passes the rate limiter.')

# Test Cases (Screenshots)
doc.add_heading('Test Cases (Screenshots)', level=2)
doc.add_paragraph('Screenshots captured while generating normal and burst requests to observe the rate-limiting behavior.\n')
doc.add_paragraph('[INSERT SCREENSHOT 1 HERE: Redis Server Running]').alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.add_paragraph('\n[INSERT SCREENSHOT 2 HERE: Normal Traffic (API responds with "pong")]').alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.add_paragraph('\n[INSERT SCREENSHOT 3 HERE: Burst Traffic (API responds with 429 Too Many Requests)]').alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_page_break()

# Analysis
doc.add_heading('Analysis', level=2)

# 1
doc.add_paragraph('1. Explain the role of:', style='List Number')
p1 = doc.add_paragraph(style='List Bullet')
p1.add_run('Token Bucket: ').bold = True
p1.add_run('Controls the rate of API requests by maintaining a metaphorical "bucket" of tokens. Each request consumes one token. If no tokens are left, the request is dropped.')
p1 = doc.add_paragraph(style='List Bullet')
p1.add_run('Bucket Capacity: ').bold = True
p1.add_run('The maximum number of tokens the bucket can hold. It determines the maximum burst size allowed at any given time.')
p1 = doc.add_paragraph(style='List Bullet')
p1.add_run('Refill Rate: ').bold = True
p1.add_run('The rate at which tokens are added back to the bucket over time. It dictates the sustained request rate permitted.')
p1 = doc.add_paragraph(style='List Bullet')
p1.add_run('Redis: ').bold = True
p1.add_run('Acts as a centralized, high-performance, in-memory data store to keep track of the token count and timestamps atomically across all application instances.')

# 2
doc.add_paragraph('2. Explain how the Token Bucket algorithm allows or rejects requests.', style='List Number')
doc.add_paragraph('When a request arrives, the algorithm calculates how much time has elapsed since the last request and refills the bucket accordingly (up to the maximum capacity). It then checks the available token count. If there is at least 1 token available, the token is consumed and the request is allowed. If there are 0 tokens available, the request is immediately rejected with an HTTP 429 error without hitting the backend service.')

# 3
doc.add_paragraph('3. Explain why Redis is required when multiple application instances are used.', style='List Number')
doc.add_paragraph('In a distributed environment with multiple application servers, local in-memory rate limiting would fail because a client could bypass limits by routing requests to different servers (e.g., behind a load balancer). Redis provides a centralized state. By executing Lua scripts atomically on a single Redis instance, all application instances share the same consistent token bucket state for a given client, preventing race conditions.')

# 4
doc.add_paragraph('4. Compare Fixed Window, Sliding Window Log, Sliding Window Counter, and Token Bucket rate-limiting algorithms.', style='List Number')
table2 = doc.add_table(rows=1, cols=3)
try:
    table2.style = 'Table Grid'
except:
    pass
hdr = table2.rows[0].cells
hdr[0].text = 'Algorithm'
hdr[1].text = 'Pros'
hdr[2].text = 'Cons'
for cell in hdr:
    for p in cell.paragraphs:
        for r in p.runs:
            r.bold = True

algos = [
    ('Fixed Window', 'Memory efficient, simple to implement.', 'Vulnerable to double-burst traffic at the edges of the window.'),
    ('Sliding Window Log', 'Highly accurate rate limiting, no burst edge issues.', 'High memory footprint (logs every timestamp), expensive to compute.'),
    ('Sliding Window Counter', 'Smoothes out traffic bursts, memory efficient compared to Log.', 'Assumes uniform distribution within the previous window.'),
    ('Token Bucket', 'Allows controlled bursts up to capacity, memory efficient, fast.', 'Requires tuning two parameters (capacity and refill rate).')
]
for item in algos:
    row_cells = table2.add_row().cells
    row_cells[0].text = item[0]
    row_cells[1].text = item[1]
    row_cells[2].text = item[2]

# 5
doc.add_paragraph('5. Analyze the behaviour of the system during normal and burst traffic.', style='List Number')
doc.add_paragraph('During normal traffic (requests arriving at a rate less than or equal to the refill rate), tokens are replenished fast enough to accommodate incoming requests, resulting in 100% allowed requests. During burst traffic (a sudden spike of requests exceeding the refill rate), the system will serve the requests until the bucket is completely drained. Subsequent requests are immediately blocked (HTTP 429). Once the burst subsides, the bucket refills over time, gradually allowing traffic again.')

# Learning Outcomes
doc.add_heading('Learning Outcomes', level=2)
p2 = doc.add_paragraph(style='List Bullet')
p2.add_run('Implemented API Rate Limiting using the Token Bucket algorithm.')
p2 = doc.add_paragraph(style='List Bullet')
p2.add_run('Configured bucket capacity and refill rate in a Spring Boot application.')
p2 = doc.add_paragraph(style='List Bullet')
p2.add_run('Used Redis to maintain centralized, atomic rate-limit information.')
p2 = doc.add_paragraph(style='List Bullet')
p2.add_run('Handled and analyzed normal and burst API traffic.')
p2 = doc.add_paragraph(style='List Bullet')
p2.add_run('Understood the importance of centralized rate limiting in distributed web applications.')

doc.save('Rate_Limiter_Report.docx')
print("Document generated successfully.")
